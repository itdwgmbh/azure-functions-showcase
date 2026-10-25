"""Read an email's Excel attachment from Graph and import it in one SQL transaction."""

import hashlib
import json
import math
import os
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import mssql_python
from azure.core.exceptions import ResourceNotFoundError
from azure.data.tables import TableClient
from azure.identity import AzureCliCredential, ManagedIdentityCredential
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

ATTACHMENT_NAME = "attachment.xlsx"
GRAPH = "https://graph.microsoft.com/v1.0"


class InvalidWorkbook(ValueError):
    """The supplied attachment cannot be imported without losing data."""


class GraphUnavailable(Exception):
    """Graph failed for a reason a retry may resolve."""


def read_rows(content: bytes) -> list[dict]:
    """Use the first worksheet, with column names in its first row."""
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (
        BadZipFile,
        InvalidFileException,
        ParseError,
        KeyError,
        ValueError,
        OSError,
    ) as exc:
        raise InvalidWorkbook(
            "The attachment is not a readable .xlsx workbook."
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        # Ignore incorrect dimension metadata sometimes emitted by Excel exporters.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header_cells = list(next(iterator, ()))
        while header_cells and header_cells[-1].value is None:
            header_cells.pop()
        if not header_cells:
            raise InvalidWorkbook("The workbook needs a header row.")
        if any(cell.data_type in ("f", "e") for cell in header_cells):
            raise InvalidWorkbook("Column names must be plain text.")
        if any(
            not isinstance(cell.value, str) or not cell.value.strip()
            for cell in header_cells
        ):
            raise InvalidWorkbook("Every column needs a non-empty text header.")
        headers = [cell.value.strip() for cell in header_cells]
        if len(set(headers)) != len(headers):
            raise InvalidWorkbook(
                "Column names must be unique after trimming whitespace."
            )

        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.value is not None for cell in cells[len(headers) :]):
                raise InvalidWorkbook(
                    f"Row {row_number} contains data without a column header."
                )
            values = []
            for cell in cells[: len(headers)]:
                if cell.data_type in ("f", "e"):
                    raise InvalidWorkbook(
                        f"Row {row_number} contains a formula or Excel error; export values first."
                    )
                value = cell.value
                if isinstance(value, (datetime, date, time)):
                    value = value.isoformat()
                elif isinstance(value, timedelta):
                    value = value.total_seconds()
                elif isinstance(value, float) and not math.isfinite(value):
                    raise InvalidWorkbook(
                        f"Row {row_number} contains a non-finite number."
                    )
                values.append(value)
            values.extend([None] * (len(headers) - len(values)))
            rows.append({"row_number": row_number, "data": dict(zip(headers, values))})
        if not rows:
            raise InvalidWorkbook("The workbook contains no data rows.")
        return rows
    except (BadZipFile, ParseError, KeyError, IndexError) as exc:
        raise InvalidWorkbook(
            "The attachment is not a readable .xlsx workbook."
        ) from exc
    finally:
        workbook.close()


def credential(*, local: bool = False):
    """Use an explicit identity; never fall back to a password."""
    if local:
        return AzureCliCredential()
    return ManagedIdentityCredential(client_id=os.environ["AZURE_CLIENT_ID"])


def graph_get(path: str, *, raw: bool = False):
    token = credential().get_token("https://graph.microsoft.com/.default").token
    url = path if path.startswith(f"{GRAPH}/") else f"{GRAPH}{path}"
    request = Request(url, headers={"Authorization": f"Bearer {token}"})
    try:
        with urlopen(request, timeout=30) as result:
            if raw:
                return result.read()
            return json.load(result)
    except HTTPError as exc:
        if exc.code == 404:
            raise ValueError("The message or attachment was not found.") from exc
        raise GraphUnavailable(f"HTTP {exc.code}") from exc
    except (URLError, TimeoutError) as exc:
        raise GraphUnavailable(type(exc).__name__) from exc


def fetch_attachments(message_id: str):
    """Download matching attachments one at a time, following Graph pagination."""
    base = f"/users/{quote(os.environ['MAILBOX'], safe='')}/messages/{quote(message_id, safe='')}"
    path = f"{base}/attachments?$select=id,name,isInline"
    found = False
    while path:
        listing = graph_get(path)
        for item in listing.get("value", []):
            if (
                item.get("@odata.type") != "#microsoft.graph.fileAttachment"
                or item.get("isInline")
                or str(item.get("name", "")).lower() != ATTACHMENT_NAME
            ):
                continue
            content = graph_get(
                f"{base}/attachments/{quote(item['id'], safe='')}/$value", raw=True
            )
            found = True
            yield item["id"], content
        path = listing.get("@odata.nextLink")
    if not found:
        raise InvalidWorkbook(f"The message has no {ATTACHMENT_NAME} attachment.")


def connect_sql(*, local: bool = False):
    return mssql_python.connect(
        f"Server=tcp:{os.environ['SQL_SERVER']},1433;"
        f"Database={os.environ['SQL_DATABASE']};"
        "Encrypt=yes;TrustServerCertificate=no;",
        token_provider=credential(local=local),
        timeout=15,
        autocommit=True,
    )


def import_message(message_id: str):
    for attachment_id, content in fetch_attachments(message_id):
        yield import_attachment(message_id, attachment_id, content)


def import_attachment(message_id: str, attachment_id: str, content: bytes) -> dict:
    content_hash = hashlib.sha256(content).hexdigest()
    with TableClient(
        endpoint=os.environ["TABLE_ENDPOINT"],
        table_name=os.environ["PROCESSED_TABLE"],
        credential=credential(),
    ) as processed:
        try:
            receipt = processed.get_entity("attachments", content_hash)
            return {
                "contentHash": content_hash,
                "importId": receipt["ImportId"],
                "rowCount": receipt["RowCount"],
                "status": "already_processed",
            }
        except ResourceNotFoundError:
            pass

        rows = read_rows(content)
        with connect_sql() as connection:
            cursor = connection.cursor()
            cursor.execute(
                "EXEC dbo.ImportWorkbook ?, ?, ?, ?",
                content_hash,
                message_id,
                attachment_id,
                json.dumps(rows, ensure_ascii=False, allow_nan=False),
            )
            result = cursor.fetchone()

        # Only mark completion after SQL commits. SQL's hash guard makes a retry
        # safe if this write fails or two queue messages process the same file.
        processed.upsert_entity(
            {
                "PartitionKey": "attachments",
                "RowKey": content_hash,
                "ImportId": result[0],
                "RowCount": result[1],
                "CompletedAt": datetime.now(UTC),
            }
        )
    return {
        "contentHash": content_hash,
        "importId": result[0],
        "rowCount": result[1],
        "status": result[2],
    }
